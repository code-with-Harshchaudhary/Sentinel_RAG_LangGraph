"""Build a portfolio research-design paper on Langfuse observability for Sentinel RAG Ops.

Design: standard_business_brief + editorial_cover. Named overrides: cover title
30 pt navy; teal kicker; 9.25 pt tables; 8.5 pt monospace schema; 8.7 pt references.
The paper separates verified Sentinel evidence from proposed Langfuse work.
"""
from pathlib import Path
import json
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "langfuse-rag-ops-paper"
OUT_DOCX = WORK / "deliverables"
OUT_PDF = WORK / "deliverables"
QA = WORK / "qa"
for folder in (WORK, OUT_DOCX, OUT_PDF, QA):
    folder.mkdir(parents=True, exist_ok=True)

NAVY = "0B2545"
BLUE = "2E74B5"
TEAL = "0A6B70"
MUTED = "596779"
INK = "25354A"
LIGHT = "F4F6F9"
FONTDIR = Path("C:/Windows/Fonts")


def figure_font(size, bold=False):
    return ImageFont.truetype(str(FONTDIR / ("calibrib.ttf" if bold else "calibri.ttf")), size)


def arrow(draw, x1, y1, x2, y2, color="#0A6B70", width=5):
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        draw.polygon([(x2, y2), (x2 - 16 * direction, y2 - 9), (x2 - 16 * direction, y2 + 9)], fill=color)
    else:
        direction = 1 if y2 > y1 else -1
        draw.polygon([(x2, y2), (x2 - 9, y2 - 16 * direction), (x2 + 9, y2 - 16 * direction)], fill=color)


def observability_architecture():
    image = Image.new("RGB", (1560, 720), "white")
    draw = ImageDraw.Draw(image)
    boxes = [
        (30, 90, 430, 280, "SENTINEL UI + API", ["Document ingestion", "Graph and vector query"]),
        (575, 50, 1015, 220, "LIGHTRAG PIPELINE", ["Parse, extract, index", "Retrieve, assemble context"]),
        (575, 405, 1015, 575, "MODEL SERVICES", ["Embedding", "Generation and judging"]),
        (1130, 90, 1530, 280, "LANGFUSE", ["Traces and observations", "Scores and experiments"]),
        (1130, 405, 1530, 575, "DECISION SURFACES", ["Dashboards and alerts", "Release comparison"]),
    ]
    for x1, y1, x2, y2, title, lines in boxes:
        fill = "#EAF4F4" if title == "LANGFUSE" else "#F2F6FA"
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=fill, outline="#B9C9D7", width=3)
        draw.text((x1 + 22, y1 + 22), title, font=figure_font(28, True), fill="#" + NAVY)
        for index, line in enumerate(lines):
            draw.text((x1 + 22, y1 + 78 + index * 36), line, font=figure_font(25), fill="#34475C")
    arrow(draw, 430, 185, 565, 140)
    arrow(draw, 795, 220, 795, 395)
    arrow(draw, 1015, 140, 1120, 185)
    arrow(draw, 1330, 280, 1330, 395)
    draw.text((30, 635), "PROPOSED TELEMETRY BOUNDARY", font=figure_font(24, True), fill="#" + TEAL)
    draw.text((405, 635), "Trace metadata is masked before export; private document text is not logged by default.", font=figure_font(24), fill="#" + MUTED)
    image.save(WORK / "observability_architecture.png")


def evaluation_loop():
    image = Image.new("RGB", (1560, 600), "white")
    draw = ImageDraw.Draw(image)
    nodes = [
        (60, 170, 330, 355, "1  CAPTURE", ["Sample traces", "Collect feedback"]),
        (420, 60, 720, 245, "2  CURATE", ["Build dataset", "Label evidence"]),
        (840, 60, 1140, 245, "3  EXPERIMENT", ["Run variants", "Compute scores"]),
        (1230, 170, 1500, 355, "4  DECIDE", ["Compare gates", "Promote or reject"]),
        (840, 360, 1140, 545, "5  MONITOR", ["Online sampling", "Regression alerts"]),
        (420, 360, 720, 545, "6  LEARN", ["Triage failures", "Add edge cases"]),
    ]
    for x1, y1, x2, y2, title, lines in nodes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill="#F2F6FA", outline="#B9C9D7", width=3)
        draw.text((x1 + 18, y1 + 20), title, font=figure_font(27, True), fill="#" + NAVY)
        for index, line in enumerate(lines):
            draw.text((x1 + 18, y1 + 76 + index * 34), line, font=figure_font(24), fill="#34475C")
    arrow(draw, 330, 245, 410, 160)
    arrow(draw, 720, 150, 830, 150)
    arrow(draw, 1140, 150, 1220, 245)
    arrow(draw, 1365, 355, 1148, 452)
    arrow(draw, 840, 452, 730, 452)
    arrow(draw, 420, 452, 195, 365)
    image.save(WORK / "continuous_evaluation_loop.png")


PAGES = [
    ("Research framing", [
        ("h2", "Problem statement"),
        ("p", "A RAG system can return a plausible answer even when retrieval was weak, context was truncated, a model call was throttled, or the answer was not supported by the selected evidence. Application uptime alone does not expose these failures. Sentinel RAG Ops already makes document status and graph state visible, but it does not yet provide a unified record of each retrieval and generation decision."),
        ("p", "This paper proposes a Langfuse-based observability and evaluation layer for Sentinel RAG Ops. The design treats an ingestion job or user query as a trace, important pipeline steps as typed observations, and quality judgments as scores. It is a research design and implementation blueprint, not a report of a completed Langfuse deployment."),
        ("h2", "Research questions"),
        ("step", ("RQ1", "Which trace structure makes graph-and-vector RAG failures attributable to ingestion, retrieval, context assembly, or generation?")),
        ("step", ("RQ2", "Which offline and online measures provide useful evidence of retrieval quality, answer grounding, latency, cost, and reliability?")),
        ("step", ("RQ3", "How can telemetry remain diagnostically useful without copying private document content into an observability store by default?")),
        ("h2", "Contribution and scope"),
        ("p", "The contribution is a system-specific trace contract, evaluation protocol, release-gate strategy, and privacy boundary grounded in the existing Sentinel architecture and current Langfuse concepts. It does not introduce a new RAG algorithm, reproduce the LightRAG benchmark, or claim that observability alone improves answer quality."),
        ("callout", "Claim boundary: the existing project has historical evidence for one successful document workflow. All Langfuse traces, dashboards, experiments, and target metrics described here are proposed work."),
    ]),
    ("Foundations and system context", [
        ("p", "Retrieval-augmented generation combines model-based generation with information selected from an external collection [1]. LightRAG extends this pattern with graph structures, entity and relationship extraction, and dual-level retrieval [2]. Sentinel RAG Ops packages that engine behind a FastAPI service and a document, graph, and retrieval interface."),
        ("h2", "Why observability is a separate engineering layer"),
        ("p", "An application log records events, but a RAG trace preserves causality across the user request, retrieval operation, selected context, model call, and result. Langfuse organizes individual observations into traces and can group traces into sessions [3]. Its typed observations include retriever, embedding, generation, evaluator, span, and event records [4]."),
        ("table", (["Layer", "Existing Sentinel responsibility", "Proposed Langfuse evidence"], [
            ["Ingestion", "Parse, chunk, extract, persist", "Stage duration, status, counts, error class"],
            ["Retrieval", "Graph and vector evidence selection", "Mode, top-k, source IDs, scores, latency"],
            ["Generation", "Keywords, extraction, final answer", "Model, prompt version, usage, latency, error"],
            ["Evaluation", "Not yet systematic", "Dataset run, deterministic checks, judge and human scores"],
            ["Operations", "Status UI and local logs", "Versioned dashboards, release comparisons, alerts"],
        ], [1500, 3900, 3960])),
        ("h2", "Design principle"),
        ("p", "Stable semantic names should describe pipeline roles rather than model brands. Langfuse guidance notes that evaluators, dashboards, saved views, and experiments depend on consistent trace structure and meaningful inputs and outputs [5]. Model identifiers, retrieval mode, storage workspace, and software release belong in attributes rather than observation names."),
        ("callout", "Operational observability answers what happened and where. Evaluation asks whether the result was good enough. The two should share trace IDs but remain conceptually distinct."),
    ]),
    ("Proposed observability architecture", [
        ("figure", ("observability_architecture.png", "Figure 1. Proposed telemetry architecture. Langfuse is an observability and evaluation plane beside the existing RAG path; it is not in the answer-critical path.")),
        ("h2", "Placement and failure isolation"),
        ("p", "Instrumentation should surround existing LightRAG boundaries instead of replacing them. The application starts a root trace when an ingestion job or query begins. Child observations record parsing, embedding, graph extraction, retrieval, context assembly, generation, and evaluation. Telemetry export must be buffered and fail open: an unavailable observability backend should not block a document job or user answer."),
        ("h2", "Identity and correlation"),
        ("table", (["Field", "Proposed value", "Reason"], [
            ["trace name", "rag.ingest or rag.query", "Stable workload grouping"],
            ["session ID", "Conversation or batch identifier", "Multi-turn and multi-document analysis"],
            ["user ID", "Pseudonymous internal identifier", "Abuse, cost, and cohort analysis"],
            ["release", "Application build identifier", "Regression comparison"],
            ["tags", "environment, route, experiment", "Known-at-start segmentation"],
            ["metadata", "workspace hash, mode, top-k", "Run-specific diagnostic context"],
        ], [1500, 3520, 4340])),
        ("p", "Langfuse is built on OpenTelemetry and can coexist with other telemetry destinations [3]. That makes a later infrastructure correlation path possible without forcing infrastructure-level spans into every LLM trace. Only observations needed to explain the RAG decision should be retained."),
    ]),
    ("Trace specification", [
        ("h2", "One trace per unit of work"),
        ("p", "The root observation should show the reviewable input and output: document metadata and terminal status for ingestion, or user question and final answer for a query. Each model invocation remains a separate generation observation so token use, latency, and failures are attributable [5]."),
        ("table", (["Trace", "Child observation", "Type", "Minimum safe attributes"], [
            ["rag.ingest", "parse_document", "span", "format, bytes, parser, status"],
            ["rag.ingest", "chunk_document", "span", "chunk policy, count, truncation"],
            ["rag.ingest", "embed_chunks", "embedding", "model, dimensions, batch size, usage"],
            ["rag.ingest", "extract_graph", "generation", "model, prompt version, entity and edge counts"],
            ["rag.query", "retrieve_context", "retriever", "mode, top-k, source hashes, retrieval scores"],
            ["rag.query", "assemble_context", "span", "selected count, context tokens, truncation"],
            ["rag.query", "generate_answer", "generation", "model, prompt version, usage, finish reason"],
            ["rag.query", "evaluate_answer", "evaluator", "metric version, score, rationale class"],
        ], [1160, 2200, 1300, 4700])),
        ("h2", "Naming and error taxonomy"),
        ("p", "Observation names remain stable across providers. Errors are normalized into parser_error, provider_auth, model_unavailable, rate_limited, storage_error, retrieval_empty, context_overflow, generation_error, and telemetry_error. Raw provider messages may be retained only after credential and content scrubbing."),
        ("code", "trace=rag.query > retrieve_context[retriever] > assemble_context[span] > generate_answer[generation] > evaluate_answer[evaluator]"),
        ("callout", "Do not log full source chunks by default. Log source hashes, counts, score summaries, token totals, and a controlled redacted preview only when a research workspace explicitly allows it."),
    ]),
    ("Evaluation methodology", [
        ("figure", ("continuous_evaluation_loop.png", "Figure 2. Proposed continuous evaluation loop. Production failures become curated test cases; offline experiments inform release decisions; online monitoring checks generalization.")),
        ("h2", "Offline and online evidence"),
        ("p", "Langfuse datasets and experiment runs can compare the same test items across prompt, model, retriever, and code variants; resulting judgments are stored as scores [6]. Online scores can then sample real traces and feed discovered edge cases back into the offline dataset. This closes the loop without treating production traffic as ground truth."),
        ("table", (["Dimension", "Primary measure", "Evidence required"], [
            ["Retrieval", "Recall@k, MRR@k, context precision", "Expected source or graded relevant passages"],
            ["Grounding", "Faithfulness / citation support", "Answer claims aligned to retrieved evidence"],
            ["Usefulness", "Answer relevance and task success", "Expected answer, rubric, or reviewer label"],
            ["Efficiency", "p50 and p95 latency, tokens, cost", "Timestamp and generation usage attributes"],
            ["Reliability", "Success rate by normalized error class", "Terminal trace status and release"],
        ], [1550, 3050, 4760])),
        ("p", "RAGAS separates context relevance, faithfulness, and answer relevance, illustrating why a single end-to-end score is insufficient [7]. Automated judges are useful for coverage but should be calibrated against human-reviewed samples and deterministic checks."),
    ]),
    ("Experimental protocol", [
        ("h2", "Dataset construction"),
        ("p", "Begin with a public, versioned corpus rather than private resumes. Create questions across direct facts, multi-hop relationships, summaries, unanswerable requests, contradictory sources, and prompt-injection attempts. Each item stores the question, expected source IDs, optional reference answer, task class, and difficulty. Split items into development and held-out evaluation sets before tuning."),
        ("h2", "Controlled comparisons"),
        ("table", (["Factor", "Candidate levels", "Control rule"], [
            ["Retrieval mode", "naive, local, global, hybrid, mix", "Same corpus, embedding space, model, and top-k"],
            ["Context budget", "small, medium, large", "Same retrieval output and answer prompt"],
            ["Prompt version", "baseline and candidate", "Same model, temperature, and evidence"],
            ["Generation model", "baseline and candidate", "Same prompt, context, and decoding policy"],
        ], [1600, 3650, 4110])),
        ("h2", "Analysis plan"),
        ("p", "Run each deterministic configuration over the identical held-out items. Report per-item results, bootstrap confidence intervals, and paired differences rather than only averages. Segment by question class and inspect failures before selecting a winner. If model sampling is enabled, repeat runs and report variance."),
        ("step", ("Gate A", "No material regression in faithfulness or expected-source retrieval on the held-out set.")),
        ("step", ("Gate B", "Latency and cost remain within a budget established from the baseline and user experience target.")),
        ("step", ("Gate C", "Security cases do not expose hidden instructions, secrets, or unsupported source content.")),
        ("callout", "Thresholds are deliberately not fabricated here. They should be fixed after a pilot, before the final comparison is run."),
    ]),
    ("Metrics, dashboards, and diagnosis", [
        ("p", "Langfuse metrics can aggregate quality scores, cost, latency, volume, tokens, and release dimensions [8]. A useful dashboard begins with decisions and failure modes, not every available field."),
        ("table", (["Question", "Dashboard slice", "Diagnostic action"], [
            ["Are answers becoming less grounded?", "Faithfulness by release and retrieval mode", "Inspect low-score traces and source selection"],
            ["Why did latency move?", "p50/p95 by observation and model", "Separate retrieval, queue, and generation time"],
            ["Where is cost concentrated?", "Tokens and cost by route, model, workspace", "Check context size, retries, and prompt changes"],
            ["What fails in production?", "Error rate by normalized class and release", "Route to parser, provider, storage, or retrieval owner"],
            ["Does a candidate generalize?", "Offline run vs sampled online score distribution", "Compare class coverage and drift"],
        ], [2800, 3140, 3420])),
        ("h2", "Alert design"),
        ("p", "Alerts should combine a meaningful threshold, minimum sample size, and time window. A single bad trace is a debugging case; a sustained rate change may be an incident. Separate availability alerts from quality monitors so a provider outage is not misclassified as an answer-quality regression."),
        ("h2", "Trace-to-action examples"),
        ("step", ("Empty retrieval", "Verify document completion, workspace identity, index version, mode, and top-k before changing the answer prompt.")),
        ("step", ("High cost", "Compare context token growth, repeated model calls, retries, and prompt version at observation level.")),
        ("step", ("Low faithfulness", "Inspect whether evidence was irrelevant, context was truncated, or generation ignored relevant passages.")),
        ("callout", "A dashboard is not evidence of improvement. A release decision requires comparable runs, preserved configurations, and inspectable failures."),
    ]),
    ("Privacy, security, and governance", [
        ("h2", "Telemetry is another data destination"),
        ("p", "Sentinel can run locally while model calls and observability export leave the device. A trace may contain user questions, retrieved passages, prompts, answers, identifiers, and provider errors. Those fields can be as sensitive as the original document. Self-hosting changes operational control but does not remove the need for minimization, access control, backups, and retention policy."),
        ("h2", "Proposed controls"),
        ("step", ("Minimize", "Default to hashes, counts, model metadata, durations, and normalized errors; omit full chunks and prompts where they are not required.")),
        ("step", ("Mask before export", "Redact credentials, email addresses, phone numbers, document paths, and configured sensitive patterns on the client side.")),
        ("step", ("Separate environments", "Use distinct Langfuse projects or environments for development, evaluation, and production; never mix public benchmark data with private workspaces.")),
        ("step", ("Restrict and expire", "Apply least-privilege access, retention limits, deletion procedures, and key rotation appropriate to the deployment.")),
        ("step", ("Audit", "Sample traces for leakage, missing redaction, label quality, and evaluator drift before expanding capture.")),
        ("p", "Langfuse documents client-side masking as the option that prevents sensitive data from leaving the application and describes server-side ingestion masking as an additional self-hosted control [9]. For this project, client-side masking is the required first boundary."),
        ("h2", "Threat-aware evaluation"),
        ("p", "The dataset should include malicious or misleading document instructions, cross-workspace retrieval attempts, secret-like strings, and unanswerable questions. Pass criteria must inspect both the final answer and the trace payload: a safe answer paired with leaked telemetry is still a failure."),
        ("callout", "Never store real API keys or private document text merely to make a trace look complete. Diagnostic value must be earned field by field."),
    ]),
    ("Implementation roadmap", [
        ("h2", "Phase 1 - trace without changing behavior"),
        ("p", "Add the Langfuse SDK behind a disabled-by-default configuration flag. Instrument one query and one ingestion path with stable names, terminal status, timing, and safe metadata. Flush telemetry asynchronously and confirm that export failure does not affect application results."),
        ("h2", "Phase 2 - establish evaluation evidence"),
        ("p", "Create a small public dataset with expected source IDs and reviewer rubrics. Run the current configuration as the named baseline. Add deterministic retrieval and citation checks first, then calibrate model-based scores against a human-reviewed subset."),
        ("h2", "Phase 3 - operationalize decisions"),
        ("p", "Add release and prompt version attributes, dashboards for latency, cost, quality, and errors, and a documented promotion checklist. Sample online traces and route failures into the dataset rather than continuously editing metrics to match observed behavior."),
        ("table", (["Risk", "Early signal", "Mitigation"], [
            ["Trace volume or cost grows", "Unexpected span count or ingestion lag", "Filter infrastructure noise; sample noncritical traces"],
            ["Sensitive content appears", "Audit finds raw PII or source text", "Block export, strengthen client masking, delete affected data"],
            ["Evaluator drift", "Human agreement falls by class", "Version judge prompts; recalibrate and retain human review"],
            ["Instrumentation changes behavior", "Latency or failure rate rises when enabled", "Buffer export, time-box flush, fail open"],
        ], [2600, 3100, 3660])),
        ("h2", "Acceptance evidence"),
        ("p", "Completion requires trace examples for both success and failure, a privacy audit, a reproducible dataset run, baseline and candidate configurations, a short error analysis, and a release decision that cites measured tradeoffs. Screenshots alone are not sufficient."),
    ]),
    ("Limitations and conclusion", [
        ("h2", "Limitations"),
        ("p", "This design has not been executed against a live Langfuse instance. It therefore cannot report trace overhead, ingestion reliability, evaluator agreement, retrieval improvements, cost, or dashboard usefulness. Langfuse features and SDK contracts may change; implementation should follow the version installed at that time. Model-based evaluation can be biased, nondeterministic, and sensitive to the judge prompt. Human review remains necessary for calibration and high-risk cases."),
        ("p", "The protocol also does not prove that the proposed metrics capture every user need. Recall@k requires meaningful relevance labels, faithfulness does not guarantee completeness, and a low-latency answer can still be wrong. Aggregate scores can hide rare but important failures, especially privacy breaches and cross-workspace retrieval."),
        ("h2", "Conclusion"),
        ("p", "A credible RAG operations layer must connect each answer to the retrieval, context, model, version, and evaluation evidence that produced it. Langfuse provides a compatible trace, score, dataset, experiment, and metrics model for that work. The proposed design keeps the answer path owned by Sentinel and LightRAG while adding an inspectable evidence plane beside it."),
        ("p", "The next milestone is intentionally small: instrument one public-data query path, verify privacy-safe trace structure, create a baseline dataset run, and publish measured results as a later empirical report. Until then, this paper is a transparent research blueprint rather than a performance claim."),
        ("callout", "Portfolio summary: designed a privacy-aware observability and evaluation architecture for a graph-and-vector RAG system, including trace contracts, experiment methodology, operational metrics, and release evidence."),
    ]),
    ("References and evidence status", [
        ("ref", ("[1] Lewis, P. et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS 2020.", "https://arxiv.org/abs/2005.11401")),
        ("ref", ("[2] Guo, Z., Xia, L., Yu, Y., Ao, T. and Huang, C. (2024; revised 2025). LightRAG: Simple and Fast Retrieval-Augmented Generation.", "https://arxiv.org/abs/2410.05779")),
        ("ref", ("[3] Langfuse. Observability data model: observations, traces, sessions, attributes, and OpenTelemetry. Accessed 1 September 2026.", "https://langfuse.com/docs/observability/data-model")),
        ("ref", ("[4] Langfuse. Observation types. Accessed 1 September 2026.", "https://langfuse.com/docs/observability/features/observation-types")),
        ("ref", ("[5] Langfuse. What does a good trace look like? Accessed 1 September 2026.", "https://langfuse.com/docs/observability/best-practices")),
        ("ref", ("[6] Langfuse. Evaluation core concepts. Accessed 1 September 2026.", "https://langfuse.com/docs/evaluation/core-concepts")),
        ("ref", ("[7] Es, S., James, J., Espinosa-Anke, L. and Schockaert, S. (2024). RAGAS: Automated Evaluation of Retrieval Augmented Generation. EACL 2024 Demo.", "https://arxiv.org/abs/2309.15217")),
        ("ref", ("[8] Langfuse. Metrics overview. Accessed 1 September 2026.", "https://langfuse.com/docs/metrics/overview")),
        ("ref", ("[9] Langfuse. Data masking for self-hosted deployments. Accessed 1 September 2026.", "https://langfuse.com/self-hosting/security/data-masking")),
        ("h2", "Project evidence status"),
        ("p", "Existing evidence: Sentinel RAG Ops architecture and one historical functional workflow documented in the companion technical case study. Proposed evidence: every Langfuse trace, score, dataset run, dashboard, alert, and promotion gate in this paper. No Langfuse credentials, telemetry, screenshots, or invented benchmark results are included."),
        ("h2", "Attribution and paper status"),
        ("p", "Author and project owner: Harsh Chaudhary. Prepared with AI-assisted research and document production. This is an independent portfolio research design paper, not a peer-reviewed publication. Core graph-RAG functionality is attributed to LightRAG and its contributors; Langfuse capabilities are attributed to its documentation."),
    ]),
]


def set_style(doc, name, size, color=INK, before=0, after=8, line=1.167, bold=False):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line
    return style


def xml_prop(parent, tag, attrs):
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    for key, value in attrs.items():
        element.set(qn(key), str(value))
    return element


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    xml_prop(properties, "w:color", {"w:val": TEAL})
    xml_prop(properties, "w:u", {"w:val": "single"})
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


def number_definition(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = max([int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))] + [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in [("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1."), ("w:lvlJc", "left")]:
        xml_prop(level, tag, {"w:val": value})
    paragraph_props = OxmlElement("w:pPr")
    xml_prop(paragraph_props, "w:ind", {"w:left": 720, "w:hanging": 360})
    tabs = OxmlElement("w:tabs")
    xml_prop(tabs, "w:tab", {"w:val": "num", "w:pos": 720})
    paragraph_props.append(tabs)
    xml_prop(paragraph_props, "w:spacing", {"w:before": 0, "w:after": 140, "w:line": 280, "w:lineRule": "auto"})
    level.append(paragraph_props)
    abstract.append(level)
    numbering.append(abstract)

    def new_id():
        number_id = max([int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))] + [0]) + 1
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(number_id))
        xml_prop(number, "w:abstractNumId", {"w:val": abstract_id})
        override = xml_prop(number, "w:lvlOverride", {"w:ilvl": 0})
        xml_prop(override, "w:startOverride", {"w:val": 1})
        numbering.append(number)
        return number_id

    return new_id


def add_table(doc, headers, rows, widths):
    assert sum(widths) == 9360
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    properties = table._tbl.tblPr
    xml_prop(properties, "w:tblW", {"w:w": 9360, "w:type": "dxa"})
    xml_prop(properties, "w:tblInd", {"w:w": 120, "w:type": "dxa"})
    margins = xml_prop(properties, "w:tblCellMar", {})
    for side, value in [("top", 80), ("bottom", 80), ("start", 120), ("end", 120)]:
        xml_prop(margins, "w:" + side, {"w:w": value, "w:type": "dxa"})
    for column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        column.set(qn("w:w"), str(width))
    borders = xml_prop(properties, "w:tblBorders", {})
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        xml_prop(borders, "w:" + side, {"w:val": "single", "w:sz": 4, "w:color": "D6DEE7"})
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        xml_prop(row_properties, "w:cantSplit", {})
        if row_index == 0:
            xml_prop(row_properties, "w:tblHeader", {})
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Inches(width / 1440)
            xml_prop(cell._tc.get_or_add_tcPr(), "w:tcW", {"w:w": width, "w:type": "dxa"})
            if row_index == 0:
                xml_prop(cell._tc.get_or_add_tcPr(), "w:shd", {"w:fill": "F2F4F7"})
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles["Table Text"]
                for run in paragraph.runs:
                    run.bold = row_index == 0
    doc.add_paragraph("", style="Table Gap")


def build():
    observability_architecture()
    evaluation_loop()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    set_style(doc, "Normal", 11)
    set_style(doc, "Title", 30, NAVY, 0, 10, 1.05, True)
    set_style(doc, "Subtitle", 15, MUTED, 0, 12, 1.15)
    for name, size, before, after, color in [
        ("Heading 1", 16, 16, 8, BLUE),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 12, 8, 4, "1F4D78"),
    ]:
        style = set_style(doc, name, size, color, before, after, 1.1, True)
        style.paragraph_format.keep_with_next = True
    set_style(doc, "Kicker", 10, TEAL, 0, 10, 1.1, True)
    set_style(doc, "Caption", 9, MUTED, 4, 8, 1.1)
    set_style(doc, "Table Text", 9.25, INK, 0, 3, 1.08)
    set_style(doc, "Table Gap", 2, INK, 0, 3, 1)
    set_style(doc, "Small", 9.5, MUTED, 0, 7, 1.1)
    set_style(doc, "Reference", 8.7, MUTED, 0, 7, 1.08)
    set_style(doc, "Callout", 11, NAVY, 8, 8, 1.15)
    set_style(doc, "Code", 8.5, INK, 5, 8, 1.1)
    doc.styles["Code"].font.name = "Consolas"
    doc.styles["Code"]._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Consolas")
    doc.styles["Code"]._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
    set_style(doc, "List Number", 11, INK, 0, 7, 280 / 240)
    doc.styles["List Number"].paragraph_format.left_indent = Inches(0.5)
    doc.styles["List Number"].paragraph_format.first_line_indent = Inches(-0.25)
    for name in ("Header", "Footer"):
        set_style(doc, name, 8.5, MUTED, 0, 0, 1.0)

    header = section.header.paragraphs[0]
    header.style = "Header"
    header.text = "LANGFUSE RAG OPS  /  RESEARCH DESIGN PAPER"
    footer = section.footer.paragraphs[0]
    footer.style = "Footer"
    footer.paragraph_format.tab_stops.clear_all()
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
    footer.add_run("Harsh Chaudhary  |  Portfolio research\t")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    properties = doc.core_properties
    properties.title = "Langfuse for Graph-RAG Operations: A Traceable Evaluation Blueprint for Sentinel RAG Ops"
    properties.author = "Harsh Chaudhary"
    properties.subject = "Portfolio research design paper"
    properties.keywords = "Langfuse, RAGOps, observability, evaluation, LightRAG, graph RAG"
    properties.comments = "Proposed instrumentation and evaluation methodology; no deployment or benchmark claim."

    paragraph = doc.add_paragraph("PORTFOLIO RESEARCH DESIGN PAPER", style="Kicker")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(55)
    paragraph = doc.add_paragraph("Langfuse for\nGraph-RAG Operations", style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph("A Traceable Evaluation Blueprint\nfor Sentinel RAG Ops", style="Subtitle")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph("Observability architecture, trace contracts, evaluation methodology,\nand privacy-aware operational evidence", style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(20)
    paragraph = doc.add_paragraph("Harsh Chaudhary\n1 September 2026", style="Small")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(18)
    doc.add_heading("Abstract", level=2)
    doc.add_paragraph("Graph-and-vector RAG systems can fail at ingestion, retrieval, context assembly, model invocation, or evaluation while still producing plausible outputs. This paper proposes a Langfuse-based observability and evaluation architecture for Sentinel RAG Ops, a LightRAG-derived document knowledge application. It defines trace boundaries, typed observations, safe metadata, an offline-to-online evaluation loop, operational metrics, release gates, and privacy controls. The work is a design study: it contributes a testable protocol and implementation roadmap, but reports no fabricated Langfuse traces or benchmark gains.")
    paragraph = doc.add_paragraph("TRACE  /  EVALUATE  /  DIAGNOSE  /  GOVERN", style="Kicker")
    paragraph.paragraph_format.space_before = Pt(12)
    doc.add_paragraph("Paper status: independent portfolio research design; proposed Langfuse integration; not peer-reviewed.", style="Small")

    new_number_id = number_definition(doc)
    markdown = [
        "# Langfuse for Graph-RAG Operations: A Traceable Evaluation Blueprint for Sentinel RAG Ops",
        "Harsh Chaudhary | 1 September 2026",
        "Independent portfolio research design paper. Proposed Langfuse integration; not a benchmark report.",
        "## Abstract",
        "Graph-and-vector RAG systems can fail at ingestion, retrieval, context assembly, model invocation, or evaluation while still producing plausible outputs. This paper proposes a Langfuse-based observability and evaluation architecture for Sentinel RAG Ops, a LightRAG-derived document knowledge application.",
    ]

    for index, (title, blocks) in enumerate(PAGES, 1):
        kicker = doc.add_paragraph(f"{index:02d}  /  RESEARCH DESIGN", style="Kicker")
        kicker.paragraph_format.page_break_before = True
        kicker.paragraph_format.keep_with_next = True
        doc.add_heading(title, level=1)
        number_id = new_number_id()
        markdown.append("\n## " + title)
        for kind, data in blocks:
            if kind == "h2":
                doc.add_heading(data, level=2)
                markdown.append("\n### " + data)
            elif kind in ("p", "small", "callout", "code"):
                style = {"p": "Normal", "small": "Small", "callout": "Callout", "code": "Code"}[kind]
                paragraph = doc.add_paragraph(data, style=style)
                if kind in ("callout", "code"):
                    xml_prop(paragraph._p.get_or_add_pPr(), "w:shd", {"w:fill": LIGHT})
                    paragraph.paragraph_format.left_indent = Inches(0.1)
                    paragraph.paragraph_format.right_indent = Inches(0.1)
                    paragraph.paragraph_format.keep_together = True
                markdown.append(("```text\n" + data + "\n```") if kind == "code" else data)
            elif kind == "step":
                paragraph = doc.add_paragraph(style="List Number")
                paragraph.add_run(data[0] + ". ").bold = True
                paragraph.add_run(data[1])
                number_properties = xml_prop(paragraph._p.get_or_add_pPr(), "w:numPr", {})
                xml_prop(number_properties, "w:ilvl", {"w:val": 0})
                xml_prop(number_properties, "w:numId", {"w:val": number_id})
                markdown.append("- **" + data[0] + "**: " + data[1])
            elif kind == "table":
                add_table(doc, *data)
                headers, rows, _ = data
                markdown.append("| " + " | ".join(headers) + " |")
                markdown.append("| " + " | ".join(["---"] * len(headers)) + " |")
                markdown.extend("| " + " | ".join(row) + " |" for row in rows)
            elif kind == "figure":
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.keep_with_next = True
                run = paragraph.add_run()
                run.add_picture(str(WORK / data[0]), width=Inches(6.5))
                drawing = run._r.xpath(".//wp:docPr")[0]
                drawing.set("descr", data[1])
                doc.add_paragraph(data[1], style="Caption")
                markdown.append("![Illustration](" + data[0] + ")\n" + data[1])
            elif kind == "ref":
                paragraph = doc.add_paragraph(data[0], style="Reference")
                paragraph.add_run("\n")
                add_hyperlink(paragraph, data[1], data[1])
                markdown.append(data[0] + " " + data[1])

    output_docx = OUT_DOCX / "Langfuse_RAG_Ops_Research_Paper.docx"
    doc.save(output_docx)
    (WORK / "Langfuse_RAG_Ops_Research_Paper.md").write_text("\n\n".join(markdown), encoding="utf-8")
    summary = {
        "intended_pages": 12,
        "section_count": len(PAGES),
        "word_count": len(re.findall(r"\b\S+\b", " ".join(markdown))),
        "preset": "standard_business_brief",
        "header_pattern": "editorial_cover",
        "claim_status": "proposed Langfuse research design; no benchmark results",
        "named_overrides": ["cover-title-30pt-navy", "teal-kicker", "9.25pt-table-text", "8.5pt-code", "8.7pt-references"],
    }
    (QA / "build_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(f"Created {output_docx}")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    build()
